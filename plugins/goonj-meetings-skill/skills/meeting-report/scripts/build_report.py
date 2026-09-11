#!/usr/bin/env python3
"""Build a bilingual meeting-report workbook from extracted content.

    python3 build_report.py content.json [--out FILE] [--no-recalc]

The agent supplies content only. Every column header, colour, formula, dropdown
and the whole How to Use sheet live here, so they are neither regenerated nor
re-derived on each run. See SKILL.md for the JSON schema.
"""
import argparse, json, os, platform, shutil, subprocess, sys, tempfile

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
except ImportError:
    sys.exit("openpyxl is not installed.  Fix:  pip install openpyxl")

EN, HI = "Arial", "Nirmala UI"
NAVY, BAND, LIGHT = "1F3864", "D9E2F3", "F2F2F2"
AMBER, GREEN, GREY, BLUE, RED = "FFF2CC", "E2EFDA", "E7E6E6", "DEEBF7", "FCE4E4"
_s = Side(style="thin", color="BFBFBF")
BOX = Border(left=_s, right=_s, top=_s, bottom=_s)


# ---------------------------------------------------------------- LibreOffice
def find_soffice():
    """Locate LibreOffice on Linux, macOS or Windows. None if absent."""
    if os.environ.get("SOFFICE") and os.path.exists(os.environ["SOFFICE"]):
        return os.environ["SOFFICE"]
    for name in ("soffice", "libreoffice", "soffice.exe"):
        p = shutil.which(name)
        if p:
            return p
    system = platform.system()
    candidates = {
        "Darwin": ["/Applications/LibreOffice.app/Contents/MacOS/soffice",
                   os.path.expanduser("~/Applications/LibreOffice.app/Contents/MacOS/soffice")],
        "Windows": [r"C:\Program Files\LibreOffice\program\soffice.exe",
                    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"],
        "Linux": ["/usr/bin/soffice", "/usr/bin/libreoffice", "/snap/bin/libreoffice",
                  "/opt/libreoffice/program/soffice"],
    }.get(system, [])
    for c in candidates:
        if os.path.exists(c):
            return c
    return None


def recalc(path):
    """Recalculate in place so formulas carry cached values. Returns a status string."""
    exe = find_soffice()
    if not exe:
        return ("SKIPPED - LibreOffice not found. The workbook is still correct: Excel "
                "recalculates on open. Only automated verification of the counts is unavailable. "
                "Install LibreOffice, or set SOFFICE=/path/to/soffice.")
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "in.xlsx")
        shutil.copy(path, src)
        try:
            subprocess.run([exe, "--headless", "--norestore", "--convert-to", "xlsx",
                            "--outdir", os.path.join(tmp, "out"), src],
                           check=True, capture_output=True, timeout=180)
        except subprocess.TimeoutExpired:
            return "SKIPPED - LibreOffice timed out after 180s."
        except subprocess.CalledProcessError as e:
            return f"SKIPPED - LibreOffice failed: {e.stderr.decode('utf-8','replace')[:200]}"
        out = os.path.join(tmp, "out", "in.xlsx")
        if not os.path.exists(out):
            return "SKIPPED - LibreOffice produced no output."
        shutil.copy(out, path)
    return "OK"


# --------------------------------------------------------------------- layout
def _hdr(ws, row, headers, widths):
    for i, (e, h) in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=f"{e}\n{h}")
        c.font = Font(name=EN, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 34


def _put(ws, r, c, val, hindi=False, size=10, bold=False):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font = Font(name=HI if hindi else EN, size=size, bold=bold)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.border = BOX
    return cell


def _autoheight(ws, r, texts, widths, minh=30, maxh=170):
    lines = 1
    for t, w in zip(texts, widths):
        if t:
            lines = max(lines, -(-len(str(t)) // max(int(w * 0.95), 8)))
    ws.row_dimensions[r].height = max(minh, min(maxh, lines * 13 + 8))


def _sheet(wb, name, headers, widths, rows, hindi_cols, note_en, note_hi):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    n = len(headers)
    t = ws.cell(row=1, column=1, value=f"{note_en}   |   {note_hi}")
    t.font = Font(name=HI, size=9, italic=True, color="555555")
    t.alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.row_dimensions[1].height = 20
    _hdr(ws, 2, headers, widths)
    r = 3
    for row in rows:
        for i, val in enumerate(row, 1):
            cell = _put(ws, r, i, val, hindi=(i in hindi_cols))
            if i == 1:
                cell.alignment = Alignment(horizontal="center", vertical="top")
        _autoheight(ws, r, row, widths)
        if r % 2 == 0:
            for c in range(1, n + 1):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=LIGHT)
        r += 1
    ws.freeze_panes = "B3"
    if r > 3:
        ws.auto_filter.ref = f"A2:{get_column_letter(n)}{r - 1}"
    return ws, r - 1


# ------------------------------------------------------------ column schemas
A_HDR = [("Sr", "क्र."), ("Action", "कार्य"), ("Action (Hindi)", "कार्य (हिन्दी)"),
         ("Owner", "ज़िम्मेदार"), ("Raised by", "किसने कहा"), ("Due", "नियत समय"),
         ("Due basis", "समय-सीमा का आधार"), ("Priority", "प्राथमिकता"),
         ("Status", "स्थिति"), ("Confidence", "विश्वसनीयता"),
         ("Evidence from transcript", "ट्रांसक्रिप्ट से प्रमाण")]
A_W = [5, 54, 54, 24, 20, 22, 13, 11, 12, 12, 60]
A_KEYS = ["sr", "en", "hi", "owner", "raised_by", "due", "due_basis", "priority",
          "status", "confidence", "evidence"]

D_HDR = [("Sr", "क्र."), ("Decision", "निर्णय"), ("Decision (Hindi)", "निर्णय (हिन्दी)"),
         ("Decided / Deferred", "तय / टाला गया"), ("By whom", "किसने"),
         ("If deferred — pending on", "टाला गया हो तो — किस पर निर्भर"),
         ("Rationale", "कारण"), ("Rationale (Hindi)", "कारण (हिन्दी)"),
         ("Evidence from transcript", "ट्रांसक्रिप्ट से प्रमाण")]
D_W = [5, 54, 54, 15, 26, 34, 44, 44, 52]
D_KEYS = ["sr", "en", "hi", "type", "by", "pending_on", "rationale_en", "rationale_hi", "evidence"]

Q_HDR = [("Sr", "क्र."), ("Open question", "खुला प्रश्न"),
         ("Open question (Hindi)", "खुला प्रश्न (हिन्दी)"),
         ("Must be answered by", "किसे उत्तर देना है"),
         ("What it blocks / impact", "क्या रुका है / प्रभाव"),
         ("Impact (Hindi)", "प्रभाव (हिन्दी)"), ("Raised in", "कहाँ उठा"), ("Status", "स्थिति")]
Q_W = [5, 52, 52, 30, 52, 52, 24, 12]
Q_KEYS = ["sr", "en", "hi", "answer_by", "impact_en", "impact_hi", "raised_in", "status"]

R_HDR = [("Sr", "क्र."), ("Risk / flag", "जोखिम"), ("Risk (Hindi)", "जोखिम (हिन्दी)"),
         ("Category", "श्रेणी"), ("Exposure in business terms", "व्यावसायिक प्रभाव"),
         ("Exposure (Hindi)", "व्यावसायिक प्रभाव (हिन्दी)"), ("Severity", "गंभीरता"),
         ("Owner", "ज़िम्मेदार"), ("Mitigation / next step", "शमन / अगला क़दम"),
         ("Evidence from transcript", "ट्रांसक्रिप्ट से प्रमाण")]
R_W = [5, 46, 46, 22, 56, 56, 11, 26, 46, 44]
R_KEYS = ["sr", "en", "hi", "category", "exposure_en", "exposure_hi", "severity",
          "owner", "mitigation", "evidence"]

C_HDR = [("Sr", "क्र."), ("Item carried forward", "पिछला लंबित विषय"),
         ("Item (Hindi)", "विषय (हिन्दी)"), ("Owner", "ज़िम्मेदार"),
         ("First raised", "पहली बार कब उठा"), ("Status this meeting", "इस बैठक में स्थिति"),
         ("Note", "टिप्पणी")]
C_W = [5, 54, 54, 24, 30, 24, 60]
C_KEYS = ["sr", "en", "hi", "owner", "first_raised", "status", "note"]

# Columns carrying Hindi CONTENT, per sheet. When language is "english" these are hidden
# rather than removed: the Summary formulas address columns by letter, so dropping any would
# silently break every count.
HINDI_COLS = {
    "Summary": ["D"],
    "Action Items": ["C"],
    "Decisions": ["C", "H"],
    "Open Questions": ["C", "F"],
    "Risks and Compliance": ["C", "F"],
    "Carry Forward": ["C"],
    "How to Use": ["D"],
}


def hide_hindi(ws):
    for col in HINDI_COLS.get(ws.title, []):
        ws.column_dimensions[col].hidden = True


STATS = [
    ("Action items — total / कुल कार्य", "=COUNTA('Action Items'!A3:A200)"),
    ("Action items — open / लंबित कार्य", "=COUNTIF('Action Items'!I3:I200,\"Open\")"),
    ("Action items — high priority / उच्च प्राथमिकता",
     "=COUNTIFS('Action Items'!H3:H200,\"High\",'Action Items'!I3:I200,\"Open\")"),
    ("Action items — no due date agreed / बिना तिथि", "=COUNTIF('Action Items'!G3:G200,\"None\")"),
    ("Decisions taken / लिए गए निर्णय", "=COUNTIF(Decisions!D3:D200,\"Decided*\")"),
    ("Decisions deferred / टाले गए निर्णय", "=COUNTIF(Decisions!D3:D200,\"Deferred*\")"),
    ("Open questions / खुले प्रश्न", "=COUNTIF('Open Questions'!H3:H200,\"Open\")"),
    ("High severity risks / उच्च जोखिम", "=COUNTIF('Risks and Compliance'!G3:G200,\"High\")"),
    ("Items carried forward still open / पुराने लंबित विषय",
     "=COUNTIF('Carry Forward'!F3:F200,\"Still open*\")"),
]

HOWTO = [
 ("Sheets in this workbook", "इस फ़ाइल की शीट्स", [
   ("Summary", "Meeting details, live counts, and the key things worth knowing if you read nothing else.",
    "बैठक का विवरण, स्वतः गिनती, और मुख्य बातें जो कुछ और न पढ़ें तो भी जान लें।"),
   ("Action Items", "Every commitment made, with owner, timing and how firm the timing actually was.",
    "हर प्रतिबद्धता, ज़िम्मेदार, समय और यह कि समय-सीमा कितनी पक्की थी।"),
   ("Decisions", "What was settled, and separately what was deferred and who it now waits on.",
    "क्या तय हुआ, और अलग से क्या टाला गया और अब किस पर निर्भर है।"),
   ("Open Questions", "Unresolved questions, each with a named person who owes an answer.",
    "अनुत्तरित प्रश्न, हर एक के साथ उत्तर देने वाले का नाम।"),
   ("Risks and Compliance", "Control gaps and data exposure, written in business terms.",
    "नियंत्रण की कमियाँ और डेटा जोखिम, व्यावसायिक भाषा में।"),
   ("Carry Forward", "Open items from earlier meetings and where each one now stands.",
    "पिछली बैठकों के लंबित विषय और उनकी वर्तमान स्थिति।")]),
 ("Column meanings", "स्तंभों का अर्थ", [
   ("Due basis — Explicit", "A date or deadline was actually stated on the call.",
    "कॉल पर तिथि या समय-सीमा स्पष्ट रूप से कही गई।"),
   ("Due basis — Relational", "Timing was tied to an event, not a date: 'before the campaign', 'by next call'.",
    "समय किसी घटना से जुड़ा था, तिथि से नहीं: 'अभियान से पहले', 'अगली कॉल तक'।"),
   ("Due basis — Conditional", "Starts only when something else happens: 'once the draft is shared'.",
    "तभी शुरू होगा जब कुछ और हो: 'मसौदा मिलने पर'।"),
   ("Due basis — None", "No timing was agreed at all. These are the ones that quietly slip — convert them to a date.",
    "कोई समय तय ही नहीं हुआ। ये ही चुपचाप छूट जाते हैं — इन्हें तिथि दें।"),
   ("Confidence — High", "The owner named themselves, or was named directly and did not object.",
    "ज़िम्मेदार ने स्वयं कहा, या सीधे नाम लिया गया और आपत्ति नहीं की।"),
   ("Confidence — Medium", "Owner inferred from context, or attributed by someone else without confirmation. Verify before acting.",
    "ज़िम्मेदार संदर्भ से निकाला गया, या किसी और ने बिना पुष्टि नाम लिया। काम से पहले जाँच लें।"),
   ("Confidence — Low", "Speaker unclear in the transcript. Confirm before circulating.",
    "ट्रांसक्रिप्ट में वक्ता स्पष्ट नहीं। साझा करने से पहले पुष्टि करें।"),
   ("Evidence", "The transcript words the row was built from, so anyone can check the reading.",
    "पंक्ति जिन शब्दों से बनी, वे ट्रांसक्रिप्ट से — ताकि कोई भी जाँच सके।")]),
 ("Working with it", "इसका उपयोग", [
   ("Update Status", "Status, Priority and Confidence are dropdowns. Changing Status updates the Summary counts automatically.",
    "Status, Priority और Confidence ड्रॉपडाउन हैं। Status बदलने पर सारांश की गिनती स्वतः बदलेगी।"),
   ("Before circulating", "Check every row marked Medium or Low confidence. Transcripts mishear names and merge speakers.",
    "Medium या Low विश्वसनीयता वाली हर पंक्ति जाँचें। ट्रांसक्रिप्ट नाम ग़लत सुनते हैं और वक्ताओं को मिला देते हैं।"),
   ("Next meeting", "The next report for this workstream carries every item still Open into its Carry Forward sheet.",
    "इस कार्यधारा की अगली रिपोर्ट हर लंबित विषय को अपनी Carry Forward शीट में ले जाएगी।"),
   ("Do not renumber", "Sr numbers are the reference across meetings. Never renumber existing rows.",
    "क्रम संख्याएँ बैठकों के बीच संदर्भ हैं। मौजूदा पंक्तियों को दोबारा क्रमांकित न करें।")]),
]


def rows_from(items, keys):
    """Dicts -> ordered tuples. Missing keys become empty strings, never KeyError."""
    return [tuple("" if it.get(k) is None else it.get(k, "") for k in keys) for it in items]


def build(content, out_path, language="bilingual"):
    english_only = language == "english"
    m = content.get("meeting", {})
    wb = Workbook()

    # ---- Summary ----
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", [3, 26, 62, 62, 14]):
        ws.column_dimensions[col].width = w

    c = ws.cell(row=2, column=2, value="MEETING REPORT / बैठक रिपोर्ट")
    c.font = Font(name=HI, bold=True, size=16, color=NAVY)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=4)
    titles = [(m.get("title_en", ""), EN)]
    if not english_only:
        titles.append((m.get("title_hi", ""), HI))
    for i, (txt, fnt) in enumerate(titles):
        c = ws.cell(row=3 + i, column=2, value=txt)
        c.font = Font(name=fnt, size=12, color="333333")
        ws.merge_cells(start_row=3 + i, start_column=2, end_row=3 + i, end_column=4)

    meta = [("Workstream / कार्यधारा", m.get("workstream", "")),
            ("Meeting type / बैठक का प्रकार", m.get("type", "")),
            ("Date / दिनांक", m.get("date", "")),
            ("Participants / प्रतिभागी", m.get("participants", "")),
            ("Recorded by / रिकॉर्ड करने वाले", m.get("recorded_by", "")),
            ("Source transcript / स्रोत", m.get("source", "")),
            ("Report language / रिपोर्ट की भाषा",
             "English only" if english_only else "English + Hindi / अंग्रेज़ी + हिन्दी")]
    r = 6
    for k, v in meta:
        a = ws.cell(row=r, column=2, value=k)
        a.font = Font(name=HI, bold=True, size=10)
        a.alignment = Alignment(vertical="top")
        a.fill = PatternFill("solid", fgColor=LIGHT)
        b = ws.cell(row=r, column=3, value=v)
        b.font = Font(name=EN, size=10)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 16 if len(str(v)) < 70 else 30
        r += 1

    r += 1
    c = ws.cell(row=r, column=2, value="AT A GLANCE / एक नज़र में")
    c.font = Font(name=HI, bold=True, size=12, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    r += 1
    for k, f in STATS:
        a = ws.cell(row=r, column=2, value=k)
        a.font = Font(name=HI, size=10)
        a.fill = PatternFill("solid", fgColor=LIGHT)
        a.border = BOX
        b = ws.cell(row=r, column=3, value=f)
        b.font = Font(name=EN, bold=True, size=11)
        b.alignment = Alignment(horizontal="left")
        b.border = BOX
        ws.row_dimensions[r].height = 16
        r += 1

    r += 1
    c = ws.cell(row=r, column=2, value="KEY POINTS / मुख्य बिंदु")
    c.font = Font(name=HI, bold=True, size=12, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    r += 1
    for i, kp in enumerate(content.get("key_points", []), 1):
        e, h = (kp + ["", ""])[:2] if isinstance(kp, list) else (kp.get("en", ""), kp.get("hi", ""))
        _put(ws, r, 2, i).alignment = Alignment(horizontal="center", vertical="top")
        _put(ws, r, 3, e)
        _put(ws, r, 4, h, hindi=True)
        _autoheight(ws, r, [e, h], [62, 62], minh=28)
        if i % 2 == 0:
            for cc in (2, 3, 4):
                ws.cell(row=r, column=cc).fill = PatternFill("solid", fgColor=LIGHT)
        r += 1
    ws.freeze_panes = "A6"

    # ---- Action Items ----
    ws_a, last_a = _sheet(wb, "Action Items", A_HDR, A_W,
        rows_from(content.get("actions", []), A_KEYS), {3},
        "One row per commitment. Edit Status as work progresses — the Summary counts update automatically.",
        "हर प्रतिबद्धता के लिए एक पंक्ति। काम बढ़ने पर Status बदलें — सारांश की गिनती स्वतः अपडेट होगी।")
    for formula, col in (('"Open,In Progress,Closed,Superseded,Dropped"', "I"),
                         ('"High,Medium,Low"', "H"), ('"High,Medium,Low"', "J"),
                         ('"Explicit,Relational,Conditional,None"', "G")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws_a.add_data_validation(dv)
        dv.add(f"{col}3:{col}200")
    if last_a >= 3:
        rng = f"I3:I{last_a}"
        for val, colour in (("Open", AMBER), ("In Progress", BLUE), ("Closed", GREEN),
                            ("Superseded", GREY)):
            ws_a.conditional_formatting.add(rng, CellIsRule(operator="equal",
                formula=[f'"{val}"'], fill=PatternFill("solid", bgColor=colour)))
        ws_a.conditional_formatting.add(f"H3:H{last_a}", CellIsRule(operator="equal",
            formula=['"High"'], font=Font(name=EN, bold=True, color="C00000")))
        for val, colour in (("Medium", AMBER), ("Low", RED)):
            ws_a.conditional_formatting.add(f"J3:J{last_a}", CellIsRule(operator="equal",
                formula=[f'"{val}"'], fill=PatternFill("solid", bgColor=colour)))
        ws_a.conditional_formatting.add(f"G3:G{last_a}", CellIsRule(operator="equal",
            formula=['"None"'], fill=PatternFill("solid", bgColor=RED)))

    # ---- Decisions ----
    ws_d, last_d = _sheet(wb, "Decisions", D_HDR, D_W,
        rows_from(content.get("decisions", []), D_KEYS), {3, 8},
        "Separates what was actually settled from what was left open. A deferred decision names who it now waits on.",
        "जो वास्तव में तय हुआ और जो खुला रह गया, दोनों अलग। टाले गए निर्णय के साथ यह भी कि अब वह किस पर निर्भर है।")
    if last_d >= 3:
        for val, colour in (("Decided", GREEN), ("Deferred", AMBER)):
            ws_d.conditional_formatting.add(f"D3:D{last_d}", FormulaRule(
                formula=[f'ISNUMBER(SEARCH("{val}",$D3))'],
                fill=PatternFill("solid", bgColor=colour)))

    # ---- Open Questions ----
    ws_q, last_q = _sheet(wb, "Open Questions", Q_HDR, Q_W,
        rows_from(content.get("questions", []), Q_KEYS), {3, 6},
        "Every question carries a named person who has to answer it. An unowned question is not tracked — it is escalated.",
        "हर प्रश्न के साथ उत्तर देने वाले का नाम है। बिना ज़िम्मेदार वाला प्रश्न ट्रैक नहीं होता — वह ऊपर भेजा जाता है।")
    dvq = DataValidation(type="list", formula1='"Open,Answered,Escalated,Dropped"', allow_blank=True)
    ws_q.add_data_validation(dvq)
    dvq.add("H3:H200")
    if last_q >= 3:
        for val, colour in (("Open", AMBER), ("Answered", GREEN)):
            ws_q.conditional_formatting.add(f"H3:H{last_q}", CellIsRule(operator="equal",
                formula=[f'"{val}"'], fill=PatternFill("solid", bgColor=colour)))

    # ---- Risks and Compliance ----
    ws_r, last_r = _sheet(wb, "Risks and Compliance", R_HDR, R_W,
        rows_from(content.get("risks", []), R_KEYS), {3, 6},
        "Exposure is written in business terms, not system terms — this is the column leadership reads.",
        "प्रभाव व्यावसायिक भाषा में लिखा गया है, तकनीकी भाषा में नहीं — नेतृत्व यही स्तंभ पढ़ता है।")
    if last_r >= 3:
        ws_r.conditional_formatting.add(f"G3:G{last_r}", CellIsRule(operator="equal",
            formula=['"High"'], fill=PatternFill("solid", bgColor=RED),
            font=Font(name=EN, bold=True, color="C00000")))
        ws_r.conditional_formatting.add(f"G3:G{last_r}", CellIsRule(operator="equal",
            formula=['"Medium"'], fill=PatternFill("solid", bgColor=AMBER)))

    # ---- Carry Forward ----
    ws_c, last_c = _sheet(wb, "Carry Forward", C_HDR, C_W,
        rows_from(content.get("carry_forward", []), C_KEYS), {3},
        "Open items from earlier meetings, reconciled against what was said this time. This is what stops items quietly disappearing.",
        "पिछली बैठकों के लंबित विषय, इस बार की चर्चा से मिलान करके। इसी से विषय चुपचाप ग़ायब होने से बचते हैं।")
    if last_c >= 3:
        for needle, colour in (("Still open", AMBER), ("Closed", GREEN)):
            ws_c.conditional_formatting.add(f"F3:F{last_c}", FormulaRule(
                formula=[f'ISNUMBER(SEARCH("{needle}",$F3))'],
                fill=PatternFill("solid", bgColor=colour)))

    # ---- How to Use ----
    ws = wb.create_sheet("How to Use")
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCD", [3, 30, 58, 58]):
        ws.column_dimensions[col].width = w
    c = ws.cell(row=2, column=2, value="HOW TO USE THIS REPORT / इस रिपोर्ट का उपयोग कैसे करें")
    c.font = Font(name=HI, bold=True, size=14, color=NAVY)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=4)
    r = 4
    for t_en, t_hi, items in HOWTO:
        c = ws.cell(row=r, column=2, value=f"{t_en} / {t_hi}")
        c.font = Font(name=HI, bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 18
        r += 1
        for a, b, cc in items:
            x = _put(ws, r, 2, a)
            x.font = Font(name=EN, bold=True, size=10)
            x.fill = PatternFill("solid", fgColor=BAND)
            _put(ws, r, 3, b)
            _put(ws, r, 4, cc, hindi=True)
            _autoheight(ws, r, [b, cc], [58, 58], minh=26)
            r += 1
        r += 1

    if english_only:
        for sheet in wb.worksheets:
            hide_hindi(sheet)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    return {"actions": last_a - 2, "decisions": last_d - 2, "questions": last_q - 2,
            "risks": last_r - 2, "carry_forward": last_c - 2}


def main():
    ap = argparse.ArgumentParser(description="Build a bilingual meeting-report workbook.")
    ap.add_argument("content", help="JSON file of extracted content")
    ap.add_argument("--out", help="output .xlsx (default: beside the JSON, named from the meeting)")
    ap.add_argument("--language", choices=["bilingual", "english"],
                    help="bilingual (default) or english. English-only hides the Hindi columns. "
                         "Overrides a \"language\" key in the content JSON.")
    ap.add_argument("--no-recalc", action="store_true", help="skip the LibreOffice recalculation")
    a = ap.parse_args()

    try:
        content = json.load(open(a.content, encoding="utf-8"))
    except json.JSONDecodeError as e:
        sys.exit(f"Content JSON is invalid: {e}")

    out = a.out
    if not out:
        m = content.get("meeting", {})
        ws_name = str(m.get("workstream", "Meeting")).replace("/", "-").replace(" ", "_")
        out = os.path.join(os.path.dirname(os.path.abspath(a.content)),
                           f"Meeting_Report_{ws_name}_{m.get('date_iso', 'undated')}.xlsx")

    language = a.language or content.get("language") or "bilingual"
    if language not in ("bilingual", "english"):
        sys.exit(f"Unknown language {language!r}. Use 'bilingual' or 'english'.")
    counts = build(content, out, language)
    status = "SKIPPED - --no-recalc" if a.no_recalc else recalc(out)
    print(json.dumps({"output": out, "language": language, "rows": counts,
                      "recalc": status}, indent=2))
    if status.startswith("SKIPPED") and not a.no_recalc:
        print(f"\nWARNING: {status}", file=sys.stderr)


if __name__ == "__main__":
    main()
